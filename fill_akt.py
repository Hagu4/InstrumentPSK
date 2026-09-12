import docx

def replace_paragraph_text(p, new_text):
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)

doc = docx.Document(r'D:\Studies\Lab and pars\Web-ProgDJ(2)\DjangoWebProject1\Akt_vnedrenia.docx')

for p in doc.paragraphs:
    text = p.text
    if text.startswith('на тему'):
        replace_paragraph_text(p, 'на тему: Разработка веб-сайта по продаже и сервисному обслуживанию бензо- и электроинструментов для ИП Громыко И.Е.')
    elif text.startswith('Автор:'):
        replace_paragraph_text(p, 'Автор: Громыко Владислав Иванович')
    elif text.startswith('Научный руководитель:'):
        replace_paragraph_text(p, 'Научный руководитель: Вертешев А.С.')
    elif text.startswith('Сущность внедряемой разработки:'):
        replace_paragraph_text(p, 'Сущность внедряемой разработки: Информационная система (веб-сайт), автоматизирующая процессы розничной продажи бензо- и электроинструментов, оформления клиентских заказов, управления товарным каталогом, а также прием и диспетчеризацию заявок на сервисное обслуживание и ремонт.')
    elif text.startswith('Форма внедрения разработки:'):
        replace_paragraph_text(p, 'Форма внедрения разработки: Развертывание клиент-серверного веб-приложения и интеграция панели администрирования в повседневные бизнес-процессы индивидуального предпринимателя.')
    elif text.startswith('Эффективность внедрения разработки'):
        replace_paragraph_text(p, 'Эффективность внедрения разработки: Сокращение трудозатрат персонала за счет автоматизации рутинных операций, привлечение новых онлайн-покупателей, а также суммарный экономический эффект в размере 351 600 рублей в год со сроком окупаемости проекта 10 месяцев.')
    elif text.startswith('Дата внедрения разработки'):
        replace_paragraph_text(p, 'Дата внедрения разработки: Июль 2026 г.')
    elif text.startswith('Предложения, замечания организации'):
        replace_paragraph_text(p, 'Предложения, замечания организации, осуществляющей внедрение: Разработка полностью соответствует техническому заданию, функциональна, удобна в эксплуатации и готова к промышленному использованию. Замечаний нет.')
    elif text.startswith('Ответственный за внедрение'):
        replace_paragraph_text(p, 'Ответственный за внедрение ____________________    Громыко И.Е.')
    elif text.strip() == '_________________________________________________________________________':
        # Remove empty line placeholders
        replace_paragraph_text(p, '')
    elif text.strip() == '____________________________________________________________________________________________________________________________________________________':
        replace_paragraph_text(p, '')
        
doc.save(r'D:\Studies\Lab and pars\Web-ProgDJ(2)\DjangoWebProject1\Akt_vnedrenia_final.docx')
print("Saved Akt_vnedrenia_final.docx")
