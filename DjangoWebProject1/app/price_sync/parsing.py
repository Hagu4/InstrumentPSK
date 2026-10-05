from dataclasses import dataclass, replace
from decimal import Decimal
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from django.conf import settings
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from .normalization import normalize_sku
from .pricing import validate_price


REQUIRED_COLUMNS = ("BRAND", "NAIMEN", "ARTIKUL", "REKOMEND_CENA")


class PriceImportValidationError(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(f"{code}: {message}")


@dataclass(frozen=True)
class SupplierRow:
    row_number: int
    brand: str
    title: str
    sku: str
    recommended_price: Decimal
    diagnostic_code: str = ""


def _file_size(file_obj) -> int:
    size = getattr(file_obj, "size", None)
    if size is not None:
        return int(size)
    current = file_obj.tell()
    file_obj.seek(0, 2)
    size = file_obj.tell()
    file_obj.seek(current)
    return size


def validate_xlsx_container(file_obj) -> None:
    filename = getattr(file_obj, "name", "")
    if filename and Path(filename).suffix.casefold() != ".xlsx":
        raise PriceImportValidationError("invalid_file_type", "Разрешены только файлы XLSX")
    if _file_size(file_obj) > settings.PRICE_IMPORT_MAX_BYTES:
        raise PriceImportValidationError("file_too_large", "Файл превышает допустимый размер")
    try:
        file_obj.seek(0)
        with ZipFile(file_obj) as archive:
            total_size = sum(item.file_size for item in archive.infolist())
            if total_size > settings.PRICE_IMPORT_MAX_UNCOMPRESSED_BYTES:
                raise PriceImportValidationError(
                    "xlsx_expands_too_large",
                    "Распакованное содержимое XLSX слишком велико",
                )
            if "[Content_Types].xml" not in archive.namelist():
                raise PriceImportValidationError("invalid_xlsx", "Файл не является XLSX")
    except (BadZipFile, OSError):
        raise PriceImportValidationError("invalid_xlsx", "Файл не является XLSX")
    finally:
        file_obj.seek(0)


def _text(value, max_length: int, field_name: str, required: bool = False) -> str:
    text = "" if value is None else str(value).strip()
    if required and not text:
        raise PriceImportValidationError(
            f"empty_{field_name}",
            f"Поле {field_name} не заполнено",
        )
    if len(text) > max_length:
        raise PriceImportValidationError(
            f"{field_name}_too_long",
            f"Поле {field_name} превышает допустимую длину",
        )
    return text


def _sku_from_cell(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, int):
        number_format = str(cell.number_format or "")
        if number_format and set(number_format) == {"0"}:
            return f"{value:0{len(number_format)}d}"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _text(value, 100, "sku")


def parse_supplier_xlsx(file_obj) -> list[SupplierRow]:
    validate_xlsx_container(file_obj)
    workbook = None
    try:
        workbook = load_workbook(
            file_obj,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        header_map = {
            str(cell.value).strip().upper(): index
            for index, cell in enumerate(header_cells)
            if cell.value is not None
        }
        missing = [column for column in REQUIRED_COLUMNS if column not in header_map]
        if missing:
            raise PriceImportValidationError(
                "missing_columns",
                "Отсутствуют обязательные колонки: " + ", ".join(missing),
            )

        parsed_rows = []
        seen_skus: dict[str, tuple[Decimal, int]] = {}
        for cells in sheet.iter_rows(min_row=2):
            if all(cell.value in (None, "") for cell in cells):
                continue
            if len(parsed_rows) >= settings.PRICE_IMPORT_MAX_ROWS:
                raise PriceImportValidationError(
                    "too_many_rows",
                    "Количество строк превышает допустимый предел",
                )
            brand = _text(cells[header_map["BRAND"]].value, 200, "brand")
            title = _text(
                cells[header_map["NAIMEN"]].value,
                500,
                "title",
                required=True,
            )
            sku = _sku_from_cell(cells[header_map["ARTIKUL"]])
            raw_price = cells[header_map["REKOMEND_CENA"]].value
            try:
                if isinstance(raw_price, str):
                    raw_price = raw_price.replace(" ", "").replace(",", ".")
                price = validate_price(raw_price)
            except ValidationError:
                raise PriceImportValidationError(
                    "invalid_price",
                    f"В строке {cells[0].row} указана некорректная цена",
                )

            supplier_row = SupplierRow(
                row_number=cells[0].row,
                brand=brand,
                title=title,
                sku=sku,
                recommended_price=price,
            )
            normalized_sku = normalize_sku(sku)
            if normalized_sku:
                previous = seen_skus.get(normalized_sku)
                if previous and previous[0] != price:
                    raise PriceImportValidationError(
                        "duplicate_conflicting_price",
                        "Один артикул содержит разные цены",
                    )
                if previous:
                    supplier_row = replace(
                        supplier_row,
                        diagnostic_code="duplicate_same_price",
                    )
                else:
                    seen_skus[normalized_sku] = (price, cells[0].row)
            parsed_rows.append(supplier_row)
        return parsed_rows
    except PriceImportValidationError:
        raise
    except (BadZipFile, KeyError, OSError, ValueError):
        raise PriceImportValidationError("invalid_xlsx", "Не удалось прочитать XLSX")
    finally:
        if workbook is not None:
            workbook.close()
        file_obj.seek(0)
